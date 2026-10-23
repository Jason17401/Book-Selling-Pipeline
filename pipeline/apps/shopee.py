"""Selling the products (sets of books - later single books too) on Shopee (蝦皮購物) as an individual seller in Taiwan.

Shopee's API is only open to Shopee Mall sellers and registered companies in Taiwan, so this uses Shopee's own bulk
upload (大量上傳) instead - no robot ever logs into Seller Center:

  Review window -> Sell products -> 'Make Shopee upload file'   (or: python -m pipeline shopee)
    1. reads any Shopee files waiting in data/listings/shopee/import (the order report, and/or a product list) and
       updates what was listed / ordered / sold ('Read Shopee report...' reads a file you choose);
    2. every product that is 'ready' (all its books checked) gets its photos (data/listings/photos/<product code>/
       01.jpg ..., put online on ImgBB when IMGBB_API_KEY is set) and a row in ONE upload file:
       data/listings/shopee/uploads/shopee_upload_<date-time>.xlsx, filled into Shopee's own template (the Excel file
       in data/listings/shopee/template/).
  YOU then: Seller Center -> 我的商品 -> 新增商品 -> 大量上傳 -> upload that file. The products wait under 尚未刊登 (not
  yet published): add the photos if they have no links, check them, publish.

Each product's code (a set: batch-set_id, e.g. 20261008-153012-S01) is its parent SKU (主商品貨號) and is in its title,
so every Shopee file that mentions a product can be matched back to its books.
"""
from __future__ import annotations

import csv
import io
import re
import shutil
import time
from pathlib import Path
from typing import Optional

from ..core import store
from ..core.config import Config
from .photohost import PhotoHostError
from .products import (ACTIVE_LISTING, apply_status, books_of, make_product_images, now_text, refresh_products)

PLATFORM = "shopee"
ON, OFF = "開啟", "關閉"           # the template's dropdown values for each delivery option
CONDITION = "二手"


class ShopeeError(RuntimeError):
    """Something you need to fix first (shown as one plain message)."""


# ---- Shopee's bulk-upload template ----------------------------------------------------------------
# (field, words that identify its column in the template's header - Chinese, English or Shopee's own keys)
COLUMN_WORDS = [
    ("min_purchase", ("最低購買數量", "minimum purchase", "minimum_purchase", "min_purchase")),
    ("cover_image", ("主商品圖片", "cover image", "cover_image")),
    ("parent_sku", ("主商品貨號", "parent sku", "parent_sku", "sku_parent")),
    ("skip", ("規格", "variation", "選項", "option")),        # variations: a product is one thing, left empty
    # weight and parcel size: left empty (optional; we do not have them)
    ("size", ("重量", "weight", "長度", "length", "寬度", "width", "高度", "height", "尺寸表", "size chart")),
    ("image", ("商品圖片", "item image", "item_image", "image")),
    ("category", ("分類", "category")),
    ("name", ("商品名稱", "product name", "product_name")),
    ("description", ("商品描述", "description")),
    ("price", ("價格", "price")),
    ("stock", ("庫存", "stock")),
    ("days_to_ship", ("備貨", "出貨天數", "days to ship", "days_to_ship", "pre-order", "preorder", "pre_order")),
    ("condition", ("保存狀況", "商品狀況", "condition")),
]
DELIVERY_WORDS = ("7-eleven", "7-11", "全家", "萊爾富", "ok mart", "ok便利", "ok超商", "店到店", "店到家", "宅配", "黑貓",
                  "新竹", "郵局", "郵政", "嘉里", "隔日到貨", "物流", "logistic", "channel")


def _whole(v, default: int = 1) -> int:
    """A stock / minimum purchase quantity as Shopee wants it: a whole number (1 when not set)."""
    try:
        return max(1, int(float(str(v).replace(",", ""))))
    except (TypeError, ValueError):
        return default


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def column_field(header: str) -> str:
    """Which field a template column is for ('' = none we fill)."""
    h = header.lower()
    if not h:
        return ""
    for field, words in COLUMN_WORDS:
        if any(w in h for w in words):
            return field
    if any(w in h for w in DELIVERY_WORDS) or h == "ok" or h.startswith("ok "):
        return "delivery"
    return ""


def _is_name_header(text: str) -> bool:
    """The product-name column's header itself (商品名稱 / Product Name / Shopee's key ps_product_name|1|0) - not an
    explanation row that merely mentions 商品名稱."""
    t = text.strip().lower()
    return t in ("商品名稱", "product name", "product_name") or t.startswith(("ps_product_name", "product_name|"))


def _layout(ws):
    """(column texts {col: header text}, first data row) of the upload sheet. The header is every row (of the first
    12) with a 商品名稱 / Product Name / product_name cell - Shopee's templates have a row of their own keys and a row
    of names - and data starts at the first empty row after the explanation rows below them."""
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=12, values_only=True)]
    heads = [i for i, r in enumerate(rows) if any(_is_name_header(_text(c)) for c in r)]
    if not heads:
        return None, None
    cols = {}
    width = max(len(r) for r in rows)
    for j in range(width):
        text = " ".join(_text(rows[i][j]) for i in heads if j < len(rows[i]) and _text(rows[i][j]))
        if text:
            cols[j + 1] = text
    last = heads[-1]
    start = next((i for i in range(last + 1, len(rows)) if not any(_text(c) for c in rows[i])), last + 4)
    return cols, start + 1


def open_workbook(path, **kw):
    """openpyxl.load_workbook that also opens Shopee's own files: their sheets say activePane="bottom_left" where the
    file format says "bottomLeft", which openpyxl refuses (Excel doesn't mind). Fixed in memory; the file is untouched."""
    import openpyxl
    import zipfile
    fixes = {b'"bottom_left"': b'"bottomLeft"', b'"bottom_right"': b'"bottomRight"', b'"top_left"': b'"topLeft"',
             b'"top_right"': b'"topRight"'}
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                for bad, good in fixes.items():
                    data = data.replace(b'activePane=' + bad, b'activePane=' + good)
            zout.writestr(info, data)
    buf.seek(0)
    return openpyxl.load_workbook(buf, **kw)


def _upload_sheet(wb):
    """(sheet, columns, first data row) of the sheet to fill: Shopee's 上傳模板 (upload template) - never its example
    sheet 參考範例, which has the same headers."""
    best = None
    for ws in wb.worksheets:
        name = ws.title.lower()
        if any(w in name for w in ("範例", "example", "sample", "說明", "guide")):
            continue
        cols, start = _layout(ws)
        if not cols:
            continue
        rank = (any(w in name for w in ("上傳", "upload", "template", "模板")), len(cols))
        if best is None or rank > best[0]:
            best = (rank, ws, cols, start)
    return best[1:] if best else None


def find_template(cfg: Config) -> Optional[Path]:
    """Shopee's template: the only (or first, by name) Excel file in data/listings/shopee/template/ - any file name,
    as Shopee named it. SHOPEE_TEMPLATE in .env may point to another file or folder."""
    p = cfg.template_path
    if p.is_file():
        files = [p]
    elif p.is_dir():
        files = sorted((f for f in p.iterdir() if f.is_file() and f.suffix.lower() in (".xlsx", ".xls")
                        and not f.name.startswith(("~$", "."))), key=lambda f: f.name.lower())
    else:
        files = []
    old = cfg.shopee_dir / "template.xlsx"                   # where older versions wanted it
    if not files and old.is_file():
        files = [old]
    xlsx = [f for f in files if f.suffix.lower() == ".xlsx"]
    if xlsx:
        return xlsx[0]
    if files:
        raise ShopeeError(f"{files[0].name} is an old .xls file: open it in Excel, File > Save As > Excel Workbook "
                          f"(.xlsx), into the same folder ({files[0].parent}), and delete the .xls")
    return None


def find_category(cfg: Config, wb) -> tuple:
    """(category id, candidates) - SHOPEE_CATEGORY_ID, or the one category in the template's category list whose name
    contains SHOPEE_CATEGORY_NAME (default 其他書籍 = Books & Magazines > Books > Other Books, 101573)."""
    if cfg.shopee_category_id:
        return cfg.shopee_category_id, []
    found = {}
    word = cfg.shopee_category_name
    if wb is not None and word:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                texts = [_text(c) for c in row]
                if any(word in t for t in texts):
                    ids = [t for t in texts if re.fullmatch(r"\d{3,}", t)]
                    if ids:
                        found[ids[0]] = " > ".join(t for t in texts if t and not t.isdigit())
    if len(found) == 1:
        return next(iter(found)), []
    return "", [f"{k}  {v}" for k, v in found.items()]


def check_template(cfg: Config) -> Path:
    """Shopee's own bulk-upload template, checked. Without it nothing is made: Shopee accepts only its own file."""
    template = find_template(cfg)
    if template is None:
        raise ShopeeError(f"No Shopee template yet, so no upload file was made and nothing was changed.\n"
                          f"Seller Center > 我的商品 > 新增商品 > 大量上傳 > 下載基本模板 (download the basic template), "
                          f"put the file (any name) in {cfg.template_path} and try again.")
    try:
        wb = open_workbook(template)
    except Exception as exc:
        raise ShopeeError(f"{template.name} could not be opened ({exc}) - download Shopee's template again")
    try:
        if _upload_sheet(wb) is None:
            raise ShopeeError(f"{template.name}: no sheet with a 商品名稱 (product name) column - is it Shopee's "
                              "bulk-upload template? Nothing was made.")
    finally:
        wb.close()
    return template


def fill_upload(cfg: Config, items: list, out: Path, template: Path, category: str) -> list:
    """Write the products into Shopee's template (one row each). Returns notes about template columns we could not
    fill or did not find. Weight and parcel size (重量 / 長度 / 寬度 / 高度) are left empty: optional, and we do not
    have them."""
    from openpyxl.styles import Alignment
    notes = []
    wb = open_workbook(template)
    found = _upload_sheet(wb)
    if found is None:
        raise ShopeeError(f"{template.name}: no sheet with a 商品名稱 (product name) column - is it Shopee's "
                          "bulk-upload template?")
    ws, cols, start = found
    fields = {j: column_field(h) for j, h in cols.items()}
    present = set(fields.values())
    for want, what in (("cover_image", "main photo"), ("parent_sku", "parent SKU / product code"),
                       ("category", "category")):
        if want not in present:
            notes.append(f"The template has no {what} column")
    for need, what in (("name", "product name"), ("description", "description"), ("price", "price")):
        if need not in present:
            raise ShopeeError(f"{template.name}: no {what} column found - is it Shopee's bulk-upload template?")
    if "condition" not in present:
        notes.append("The template has no condition column: Shopee keeps its default - the description says 二手")
    from openpyxl.utils import get_column_letter
    for j, f in fields.items():                 # wide enough to read the title and description in Excel
        if f in ("name", "description"):
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, 60)
    for i, it in enumerate(items):
        r = start + i
        images = list(it["image_urls"])
        for j, f in sorted(fields.items()):
            value = None
            if f == "category":
                value = (int(category) if category.isdigit() else category) or None
            elif f == "name":
                value = it["title"]
            elif f == "description":
                value = it["description"].replace("\r\n", "\n")         # every line, as in Sell products
            elif f == "parent_sku":
                value = it["code"]
            elif f == "price":
                value = int(float(it["price"]))
            elif f == "stock":
                value = _whole(it.get("stock"))
            elif f == "min_purchase":
                value = _whole(it.get("min_purchase"))
            elif f == "cover_image":
                value = images.pop(0) if images else None
            elif f == "image":
                value = images.pop(0) if images else None
            elif f == "days_to_ship" and cfg.shopee_days_to_ship:
                value = int(cfg.shopee_days_to_ship) if cfg.shopee_days_to_ship.isdigit() else cfg.shopee_days_to_ship
            elif f == "condition":
                value = CONDITION
            elif f == "delivery":
                h = cols[j].lower()
                value = ON if any(w in h for w in cfg.shopee_logistics) else OFF
            if value is not None:
                cell = ws.cell(r, j, value)
                if f == "description":          # Excel shows every line of it, not only the first
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        lines = it["description"].count("\n") + 1
        ws.row_dimensions[r].height = min(409, 15 * lines)          # 409 = the most Excel allows
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return notes


# ---- photos online (optional) --------------------------------------------------------------------
def upload_images(cfg: Config, code: str, paths: list, http=None) -> list:
    """The photos' links on ImgBB (pipeline/apps/photohost.py), uploaded first. No IMGBB_API_KEY: no links (you add
    the photos by hand)."""
    from .photohost import put_online
    return put_online(cfg, code, paths, http=http)


# ---- links -------------------------------------------------------------------------------------
def parse_url(url: str) -> tuple:
    """(shop id, item id) from a Shopee product link: shopee.tw/<name>-i.<shop>.<item> or shopee.tw/product/<shop>/<item>."""
    m = re.search(r"-i\.(\d+)\.(\d+)", url or "") or re.search(r"/product/(\d+)/(\d+)", url or "")
    return (m.group(1), m.group(2)) if m else ("", "")


def item_url(shop_id: str, item_id: str) -> str:
    return f"https://shopee.tw/product/{shop_id}/{item_id}" if shop_id and item_id else ""


# ---- listings.csv helpers ----------------------------------------------------------------------
def _listing(listings: list, code: str) -> Optional[dict]:
    return next((li for li in listings if li["product_code"] == code and li["platform"] == PLATFORM), None)


def _new_listing(listings: list, code: str) -> dict:
    li = {"product_code": code, "platform": PLATFORM}
    listings.append(li)
    return li


def _other_active(listings: list, code: str) -> list:
    return [li for li in listings if li["product_code"] == code and li["platform"] != PLATFORM
            and li.get("status") in ACTIVE_LISTING]


# ---- export ------------------------------------------------------------------------------------
def _is_plain_sheet(path: str) -> bool:
    """An upload file made WITHOUT Shopee's template by an older version (one sheet called 'upload'): Shopee does not
    accept it."""
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            names = re.findall(r'<sheet [^>]*name="([^"]*)"', z.read("xl/workbook.xml").decode("utf-8", "replace"))
    except Exception:
        return False
    return names == ["upload"]


def undo_plain_exports(cfg: Config) -> list:
    """Sets an older version put in a plain sheet (no template) count as listed although Shopee never got them: they
    go back to ready (and into the next upload file). Returns notes."""
    if not cfg.listings_path.exists():
        return []
    notes = []
    with store.locked(cfg.csv_path), store.locked(cfg.products_path), store.locked(cfg.listings_path):
        listings = store.read_table(cfg.listings_path)
        bad = [li for li in listings if li["platform"] == PLATFORM and li.get("status") == "exported"
               and li.get("export_file") and _is_plain_sheet(li["export_file"])]
        if not bad:
            return []
        books = store.read_rows(cfg.csv_path)
        sets = store.read_table(cfg.products_path)
        for li in bad:
            code = li["product_code"]
            s = next((x for x in sets if x["product_code"] == code), None)
            if s is not None and s["status"] == "listed":
                apply_status(s, "ready", books_of(cfg, code, books))
                s["updated_at"] = now_text()
            listings.remove(li)
            notes.append(f"{code}: its last upload file was not Shopee's template - back to ready, made again")
        store.write_rows(cfg.csv_path, books)
        store.write_table(cfg.products_path, sets, store.PRODUCT_COLUMNS)
        store.write_table(cfg.listings_path, listings, store.LISTING_COLUMNS)
    return notes


def export(cfg: Config, redo=(), http=None) -> dict:
    """Make the upload file for every product that is ready and not on Shopee yet (and the products in `redo` again).
    Returns {"file", "made": [codes], "drafts": [(code, why)], "notes": [...], "photos": folder}."""
    redo = set(redo or ())
    res = {"file": None, "made": [], "drafts": [], "notes": [], "photos": cfg.photos_dir}
    res["notes"] += undo_plain_exports(cfg)
    sets = refresh_products(cfg)
    known = {s["product_code"] for s in sets}
    for code in redo - known:
        res["notes"].append(f"{code}: no such product")
    listings = store.read_table(cfg.listings_path)
    active = {li["product_code"] for li in listings if li["platform"] == PLATFORM and li.get("status") in ACTIVE_LISTING}
    todo = [s for s in sets if (s["status"] == "ready" and s["product_code"] not in active)
            or (s["product_code"] in redo and s["status"] not in ("draft", "sold"))]
    res["drafts"] = [(s["product_code"], s.get("errors") or "not ready") for s in sets if s["status"] == "draft"]
    for s in sets:
        if s["product_code"] in redo and s["status"] in ("draft", "sold"):
            res["notes"].append(f"{s['product_code']} is {s['status']}: not put in the file")
    if not todo:
        return res
    template = check_template(cfg)          # no template: stop here - nothing made, nothing changed
    wb = open_workbook(template)            # not read_only: Shopee's files give wrong sheet sizes
    category, candidates = find_category(cfg, wb)
    wb.close()
    if not category:                 # optional in Shopee's template: Shopee then suggests one for each set
        note = ("No category given: Shopee suggests one for each product under 尚未刊登 - check it there (or set "
                "SHOPEE_CATEGORY_ID in .env to always use one)")
        if candidates:
            note += f". Categories containing '{cfg.shopee_category_name}': " + "; ".join(candidates[:10])
        res["notes"].append(note)
    books = store.read_rows(cfg.csv_path)
    items = []
    for s in todo:
        code = s["product_code"]
        group = books_of(cfg, code, books)
        if not group:
            res["notes"].append(f"{code}: its books are not in books.csv any more - skipped")
            continue
        paths = make_product_images(s, group, cfg, cfg.photos_dir / code)
        try:
            urls = upload_images(cfg, code, paths, http)
        except PhotoHostError as exc:
            res["notes"].append(f"{code}: photos not put online ({exc}) - add them by hand when you publish")
            urls = []
        items.append({"code": code, "title": s["title"], "description": s["description"],
                      "price": s["price"], "stock": s.get("stock"), "min_purchase": s.get("min_purchase"),
                      "image_urls": urls,
                      "images": [str(p) for p in paths]})
    if not items:
        return res
    out = cfg.uploads_dir / f"shopee_upload_{time.strftime('%Y%m%d-%H%M%S')}.xlsx"
    n = 1
    while out.exists():                        # never overwrite an earlier upload file
        n += 1
        out = out.with_name(f"{out.stem.split('_v')[0]}_v{n}.xlsx")
    res["notes"] += fill_upload(cfg, items, out, template, category)
    res["file"] = out
    stamp = now_text()
    with store.locked(cfg.csv_path), store.locked(cfg.products_path), store.locked(cfg.listings_path):
        books = store.read_rows(cfg.csv_path)
        sets = store.read_table(cfg.products_path)
        listings = store.read_table(cfg.listings_path)
        for it in items:
            code = it["code"]
            s = next(x for x in sets if x["product_code"] == code)
            if s["status"] not in ("listed", "reserved", "sold"):
                apply_status(s, "listed", books_of(cfg, code, books))
                s["errors"] = ""
            s["images"] = ";".join(it["images"])
            s["updated_at"] = stamp
            li = _listing(listings, code) or _new_listing(listings, code)
            for k in store.LISTING_COLUMNS[2:]:
                li[k] = ""
            li.update({"status": "exported", "title": it["title"], "price": str(int(float(it["price"]))),
                       "currency": "TWD", "export_file": str(out), "exported_at": stamp, "updated_at": stamp,
                       "shopee_sku": code, "shopee_category_id": category, "shopee_condition": CONDITION,
                       "shopee_logistics": ",".join(cfg.shopee_logistics),
                       "shopee_days_to_ship": cfg.shopee_days_to_ship,
                       "errors": "" if it["image_urls"] else "add the photos by hand when you publish"})
            res["made"].append(code)
        store.write_rows(cfg.csv_path, books)
        store.write_table(cfg.products_path, sets, store.PRODUCT_COLUMNS)
        store.write_table(cfg.listings_path, listings, store.LISTING_COLUMNS)
    return res


# ---- reading Shopee's files back -------------------------------------------------------------------
ORDER_ID = ("訂單編號", "訂單號碼", "order id", "order sn", "order_sn", "order no")
ITEM_ID = ("商品id", "商品 id", "product id", "product_id", "item id", "item_id")


def read_report(path: Path, password: str = "") -> list:
    """Every sheet of a Shopee file as lists of cell texts. Shopee's order report is password protected (the last
    6 digits of your phone, usually): with SHOPEE_EXPORT_PASSWORD and msoffcrypto-tool it is opened directly;
    otherwise open it in Excel once and save it without a password."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp950", "big5"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ShopeeError(f"{path.name}: unknown text encoding")
        return [[[_text(c) for c in row] for row in csv.reader(io.StringIO(text))]]
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            raise ShopeeError(f"{path.name} is an old .xls file: open it in Excel and save it as .xlsx")
        book = xlrd.open_workbook(str(path))
        return [[[_text(c) for c in sh.row_values(i)] for i in range(sh.nrows)] for sh in book.sheets()]
    import openpyxl
    import zipfile
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        if not password:
            raise ShopeeError(f"{path.name} could not be opened ({exc.__class__.__name__}) - if it is Shopee's "
                              "password-protected report: put the password in SHOPEE_EXPORT_PASSWORD, or open it in "
                              "Excel and save it without a password")
        try:
            import msoffcrypto
        except ImportError:
            raise ShopeeError(f"{path.name} is password protected: .venv\\Scripts\\python -m pip install "
                              "msoffcrypto-tool (or open it in Excel and save it without a password)")
        buf = io.BytesIO()
        with open(path, "rb") as f:
            office = msoffcrypto.OfficeFile(f)
            office.load_key(password=password)
            office.decrypt(buf)
        buf.seek(0)
        wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)
    sheets = [[[_text(c) for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    wb.close()
    return sheets


def _find(headers: list, words, avoid=()) -> Optional[int]:
    """The column whose header contains the first of `words` that any header contains (words in order of
    preference)."""
    low = [h.lower() for h in headers]
    for w in words:
        for i, h in enumerate(low):
            if w in h and not any(a in h for a in avoid):
                return i
    return None


def _table(rows: list):
    """(header, data rows) - the header is the row (of the first 10) that names an order id or a product id."""
    for i, r in enumerate(rows[:10]):
        if _find(r, ORDER_ID) is not None or _find(r, ITEM_ID) is not None:
            return r, rows[i + 1:]
    return None, []


def _number(text: str) -> Optional[float]:
    t = re.sub(r"[^\d.\-]", "", text or "")
    try:
        return float(t) if t not in ("", "-", ".") else None
    except ValueError:
        return None


def order_state(text: str) -> str:
    """completed / cancelled / returned / open, from Shopee's order status (訂單狀態)."""
    t = (text or "").lower()
    if any(w in t for w in ("不成立", "取消", "cancel")):
        return "cancelled"
    if any(w in t for w in ("退貨", "退款", "return", "refund")):
        return "returned"
    if any(w in t for w in ("已完成", "完成", "completed")):
        return "completed"
    return "open"


def _codes_in(cells: list, codes: list) -> Optional[str]:
    text = " ".join(cells)
    hits = [c for c in codes if c in text]
    return max(hits, key=len) if hits else None


def sync(cfg: Config, files: list) -> dict:
    """Read Shopee's order report(s) and product list(s): match every row to a product by its code, then record item ids,
    links, orders, reserved / sold. Returns {"lines": [what changed], "takedown": [...], "read": [files]}."""
    res = {"lines": [], "takedown": [], "read": [], "problems": []}
    orders, products = {}, {}
    sets_now = refresh_products(cfg)                 # every product known, even before the first upload file
    codes = sorted((s["product_code"] for s in sets_now), key=len, reverse=True)
    for f in files:
        try:
            sheets = read_report(Path(f), cfg.shopee_export_password)
        except ShopeeError as exc:
            res["problems"].append(str(exc))
            continue
        res["read"].append(f)
        for rows in sheets:
            header, data = _table(rows)
            if header is None:
                continue
            oid = _find(header, ORDER_ID)
            if oid is not None:
                col = {"status": _find(header, ("訂單狀態", "order status", "status")),
                       "price": _find(header, ("商品總價", "商品活動價格", "商品原價", "訂單小計", "product subtotal",
                                               "deal price", "price"), avoid=("運費", "shipping")),
                       "tracking": _find(header, ("包裹查詢號碼", "物流追蹤", "tracking")),
                       "created": _find(header, ("訂單成立", "order creation", "create time")),
                       "done": _find(header, ("訂單完成時間", "完成時間", "order complete", "completed time")),
                       "payout": _find(header, ("實收", "撥款金額", "入帳金額", "payout", "escrow"))}
                fee_cols = [i for i, h in enumerate(header)
                            if any(w in h.lower() for w in ("手續費", "服務費", "處理費", "fee"))]
                for r in data:
                    code = _codes_in(r, codes)
                    if not code or oid >= len(r) or not r[oid]:
                        continue

                    def cell(k):
                        i = col[k]
                        return r[i] if i is not None and i < len(r) else ""
                    fees = sum(_number(r[i]) or 0 for i in fee_cols if i < len(r))
                    orders[(code, r[oid])] = {"sn": r[oid], "status": cell("status"), "price": cell("price"),
                                              "tracking": cell("tracking"), "created": cell("created"),
                                              "done": cell("done"), "payout": cell("payout"),
                                              "fees": f"{fees:g}" if fee_cols else ""}
                continue
            iid = _find(header, ITEM_ID)
            st = _find(header, ("狀態", "status"))
            for r in data:
                code = _codes_in(r, codes)
                if code and iid < len(r) and re.fullmatch(r"\d+", r[iid] or ""):
                    products[code] = {"item_id": r[iid], "status": r[st] if st is not None and st < len(r) else ""}
    if not orders and not products:
        return res
    stamp = now_text()
    with store.locked(cfg.csv_path), store.locked(cfg.products_path), store.locked(cfg.listings_path):
        books = store.read_rows(cfg.csv_path)
        sets = store.read_table(cfg.products_path)
        listings = store.read_table(cfg.listings_path)
        by_code = {s["product_code"]: s for s in sets}
        for code, p in products.items():
            s = by_code[code]
            li = _listing(listings, code) or _new_listing(listings, code)
            li["shopee_item_id"] = p["item_id"]
            li["shopee_sku"] = code
            li["url"] = li.get("url") or item_url(cfg.shopee_shop_id, p["item_id"])
            li["last_synced_at"] = stamp
            t = p["status"].lower()
            gone = any(w in t for w in ("下架", "unlist", "delist", "hidden"))
            banned = any(w in t for w in ("違規", "禁止", "banned", "刪除", "deleted"))
            if li.get("status") == "sold":
                continue
            if gone or banned:
                if li.get("status") != ("rejected" if banned else "delisted"):
                    li["status"], li["delisted_at"] = ("rejected" if banned else "delisted"), stamp
                    if s["status"] in ("listed", "ready"):
                        apply_status(s, "withdrawn", books_of(cfg, code, books))
                        s["errors"] = (f"{'removed by Shopee' if banned else 'taken off Shopee'} - to sell it again: "
                                       f"python -m pipeline product-status {code} ready")
                    res["lines"].append(f"{code}: {'removed by Shopee' if banned else 'not on sale on Shopee'} "
                                        f"({p['status']})")
                continue
            if li.get("status") != "live":
                li["status"], li["listed_at"] = "live", li.get("listed_at") or stamp
                res["lines"].append(f"{code}: on Shopee, item {p['item_id']}")
            if s["status"] in ("ready", "draft"):
                apply_status(s, "listed", books_of(cfg, code, books))
        rank = {"completed": 3, "returned": 2, "open": 1, "cancelled": 0}
        best = {}
        for (code, sn), o in orders.items():
            o["state"] = order_state(o["status"])
            if code not in best or rank[o["state"]] > rank[best[code]["state"]]:
                best[code] = o
        for code, o in best.items():
            s = by_code[code]
            group = books_of(cfg, code, books)
            li = _listing(listings, code) or _new_listing(listings, code)
            li["last_synced_at"] = stamp
            li["shopee_sku"] = code
            state = o["state"]
            if state == "cancelled":
                if li.get("shopee_order_sn") == o["sn"] and s["status"] == "reserved":
                    apply_status(s, "listed", group)
                    li["status"] = "live"
                    li["shopee_order_sn"], li["shopee_order_status"] = "", f"{o['sn']} {o['status']}"
                    s["errors"] = ""
                    res["lines"].append(f"{code}: order {o['sn']} cancelled - on sale again")
                continue
            before = (s["status"], li.get("shopee_order_status"))
            li.update({"shopee_order_sn": o["sn"], "shopee_order_status": o["status"],
                       "shopee_tracking_no": o["tracking"] or li.get("shopee_tracking_no", ""),
                       "shopee_fees": o["fees"] or li.get("shopee_fees", ""),
                       "shopee_payout": o["payout"] or li.get("shopee_payout", "")})
            if state == "completed":
                li["status"] = "sold"
                if s["status"] != "sold":
                    apply_status(s, "sold", group)
                    price = _number(o["price"])
                    s.update({"sold_on": PLATFORM, "sold_at": o["done"] or o["created"] or stamp,
                              "sold_price": f"{price:g}" if price is not None else s.get("price", ""),
                              "errors": ""})
                    res["lines"].append(f"{code}: SOLD on Shopee (order {o['sn']})")
                    for other in _other_active(listings, code):
                        other["status"], other["delisted_at"] = "delisted", stamp
                        res["takedown"].append(f"{code} sold on Shopee: take it down on {other['platform']}"
                                               + (f" {other['url']}" if other.get("url") else ""))
            else:
                if li.get("status") in ("exported", "", None):
                    li["status"], li["listed_at"] = "live", li.get("listed_at") or stamp
                if s["status"] != "sold" or state == "returned":
                    apply_status(s, "reserved", group)
                s["errors"] = (f"Shopee order {o['sn']}: {o['status']} (return / refund) - when the books are back: "
                               f"python -m pipeline product-status {code} ready") if state == "returned" else ""
                if before != (s["status"], li.get("shopee_order_status")):
                    res["lines"].append(f"{code}: ordered on Shopee (order {o['sn']}, {o['status']})")
            s["updated_at"] = stamp
        store.write_rows(cfg.csv_path, books)
        store.write_table(cfg.products_path, sets, store.PRODUCT_COLUMNS)
        store.write_table(cfg.listings_path, listings, store.LISTING_COLUMNS)
    return res


def set_link(cfg: Config, code: str, url: str) -> dict:
    """You published a product and pasted its Shopee link: its item id and link are saved, and it counts as on sale."""
    shop, item = parse_url(url)
    if not item:
        raise ShopeeError("That is not a Shopee product link (it should look like shopee.tw/...-i.123.456)")
    stamp = now_text()
    with store.locked(cfg.csv_path), store.locked(cfg.products_path), store.locked(cfg.listings_path):
        books = store.read_rows(cfg.csv_path)
        sets = store.read_table(cfg.products_path)
        listings = store.read_table(cfg.listings_path)
        s = next((x for x in sets if x["product_code"] == code), None)
        if s is None:
            raise ShopeeError(f"product {code} not found")
        li = _listing(listings, code) or _new_listing(listings, code)
        li.update({"url": url.strip(), "shopee_item_id": item, "shopee_sku": code, "updated_at": stamp})
        if li.get("status") not in ("sold",):
            li["status"], li["listed_at"] = "live", li.get("listed_at") or stamp
        if s["status"] in ("ready", "draft", "withdrawn"):
            apply_status(s, "listed", books_of(cfg, code, books))
        s["updated_at"] = stamp
        store.write_rows(cfg.csv_path, books)
        store.write_table(cfg.products_path, sets, store.PRODUCT_COLUMNS)
        store.write_table(cfg.listings_path, listings, store.LISTING_COLUMNS)
    return dict(li)


# ---- the whole thing (the Sell window's buttons, or `python -m pipeline shopee`) ---------------------------------------------------------------
def import_files(cfg: Config) -> list:
    d = cfg.shopee_import
    if not d.exists():
        return []
    return sorted(p for p in d.iterdir() if p.is_file() and p.suffix.lower() in (".xlsx", ".xls", ".csv")
                  and not p.name.startswith("~$"))


def run(cfg: Config, files=None, do_export: bool = True, redo=(), http=None) -> dict:
    """1. read Shopee's files (given, or everything in data/listings/shopee/import - moved to import/done after),
    2. make the upload file for the ready products. Returns {"sync": ..., "export": ..., "text": [lines to show]}."""
    cfg.shopee_import.mkdir(parents=True, exist_ok=True)
    from_folder = files is None
    files = import_files(cfg) if from_folder else [Path(f) for f in files]
    out = {"sync": None, "export": None, "text": []}
    t = out["text"]
    if files:
        res = out["sync"] = sync(cfg, files)
        t.append(f"Read {len(res['read'])} Shopee file(s):")
        t += [f"  {line}" for line in res["lines"]] or ["  nothing new for any product"]
        t += [f"  TAKE DOWN: {x}" for x in res["takedown"]]
        t += [f"  PROBLEM: {x}" for x in res["problems"]]
        if from_folder:
            done = cfg.shopee_import / "done"
            done.mkdir(exist_ok=True)
            for f in res["read"]:
                shutil.move(str(f), str(done / f"{time.strftime('%Y%m%d-%H%M%S')}_{Path(f).name}"))
    if do_export:
        res = out["export"] = export(cfg, redo=redo, http=http)
        if res["file"]:
            t.append(f"Upload file for {len(res['made'])} product(s): {res['file']}")
            t += [f"  {c}" for c in res["made"]]
            t.append(f"Photos: {res['photos']}\\<product code>"
                     + (f" (online: {cfg.photo_service})" if cfg.photo_service else ""))
            t.append("Next: Seller Center > 我的商品 > 新增商品 > 大量上傳: upload the file. The products wait under "
                     "尚未刊登: " + ("check them" if cfg.photo_service else "add each product's photos (01.jpg first)")
                     + " and publish. When an order comes in: Sell products > Read Shopee report... (the order "
                     "report from Seller Center).")
        else:
            t.append("No product is ready for a new Shopee listing.")
        t += [f"  NOTE: {n}" for n in res["notes"]]
        if res["drafts"]:
            t.append(f"{len(res['drafts'])} product(s) not ready yet:")
            t += [f"  {code}: {why}" for code, why in res["drafts"]]
    return out
