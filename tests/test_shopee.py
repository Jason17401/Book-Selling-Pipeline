"""Selling products (sets of books): data/products.csv, data/listings.csv, the Shopee upload file, and reading Shopee's reports back."""

import csv

import openpyxl
import pytest
from PIL import Image
from synth import TW_ISBNS, book_back, set_front

from pipeline.apps import shopee
from pipeline.apps.products import (
    change_product_status,
    refresh_products,
    save_product_fields,
    set_description_zh,
    set_title_zh,
)
from pipeline.core import store
from pipeline.core.config import Config

CODE = "20261008-153012-S01"
TITLES = [
    "狄更斯的耶誕頌",
    "鬼屋裡的音樂家",
    "愛爾蘭任務",
    "月光下的魔笛",
    "再窮也要玩歐洲",
    "社交高手必修課",
    "衝業績一定有效的10種態度",
    "Google必修的圖表簡報術：Google總監首度公開，教你做對圖表",
]


def make_template(path, category=True):
    """A stand-in for Shopee's bulk-upload template: a row of Shopee's keys, a row of names, three rows of
    explanations, data from row 6; and a category list."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"
    cols = [
        ("ps_category", "分類代碼"),
        ("ps_product_name", "商品名稱"),
        ("ps_product_description", "商品描述"),
        ("ps_minimum_purchase_quantity|0|0", "最低購買數量"),
        ("ps_sku_parent_short", "主商品貨號"),
        ("et_title_variation_1", "規格名稱1"),
        ("ps_price", "價格"),
        ("ps_stock", "庫存"),
        ("ps_item_cover_image", "主商品圖片"),
        ("ps_item_image_1", "商品圖片 1"),
        ("ps_item_image_2", "商品圖片 2"),
        ("ps_weight", "重量"),
        ("ps_length", "長度"),
        ("ps_width", "寬度"),
        ("ps_height", "高度"),
        ("channel_30005", "7-ELEVEN"),
        ("channel_30006", "全家"),
        ("channel_39303", "黑貓宅急便"),
        ("ps_product_pre_order_dts", "較長備貨天數"),
    ]
    for j, (key, name) in enumerate(cols, 1):
        for i, v in enumerate((key, name, "必填", "說明", "格式"), 1):
            ws.cell(i, j, v)
    if category:
        cat = wb.create_sheet("分類")
        cat.append(["分類代碼", "分類"])
        cat.append(["101573-書籍及雜誌期刊/書籍/其他書籍", "101573", "2 - 15"])  # as in Shopee's real template
        cat.append(["101541-書籍及雜誌期刊/書籍/繪本童書", "101541", "2 - 15"])
    if path.suffix.lower() != ".xlsx":  # the template folder: saved under the name Shopee gives it
        path = path / "Shopee_mass_upload_2026-10-14_basic_template.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return [name for _, name in cols]


@pytest.fixture
def cfg(tmp_path):
    c = Config(data_dir=tmp_path / "data", market_providers=(), grid=(2, 4), set_size=8, set_discount=0.0)
    c.ensure_dirs()
    d = c.processed / "20261008-153012" / "S01"
    d.mkdir(parents=True)
    front = d / "front.jpg"
    set_front(8, (2, 4)).save(front)
    rows = []
    for i in range(8):
        bc = d / f"{i + 1:02d}.jpg"
        book_back(TW_ISBNS[i], size=(300, 400)).save(bc)
        rows.append(
            {
                "sku": f"{CODE}-{i + 1:02d}",
                "batch": "20261008-153012",
                "set_id": "S01",
                "position": str(i + 1),
                "isbn13": TW_ISBNS[i],
                "title": TITLES[i],
                "author": "瑪麗．波．奧斯本",
                "publisher": "小天下",
                "condition": "as_new",
                "price": "100",
                "currency": "TWD",
                "market_price": "250",
                "market_currency": "TWD",
                "set_photo": str(front),
                "front_photo": str(front),
                "back_photo": str(bc),
                "status": "validated",
            }
        )
    rows[3]["status"] = "to_check"  # one book not checked yet
    store.write_rows(c.csv_path, rows)
    return c


def books(cfg):
    return {r["position"]: r for r in store.read_rows(cfg.csv_path)}


def check_all(cfg):
    store.update_rows(cfg.csv_path, lambda rows: [r.update(status="validated") for r in rows])


def the_set(cfg):
    return next(s for s in store.read_table(cfg.products_path) if s["product_code"] == CODE)


def the_listing(cfg):
    return next(li for li in store.read_table(cfg.listings_path) if li["product_code"] == CODE)


def test_the_set_follows_its_books(cfg):
    s = refresh_products(cfg)[0]
    assert s["product_code"] == CODE and s["status"] == "draft" and "book(s) 4 not checked yet" in s["errors"]
    assert s["book_count"] == "8" and s["price"] == "800" and s["currency"] == "TWD"
    assert "weight_g" not in s and "length_cm" not in s  # no weight or parcel size: we do not have them
    assert s["title"] == f"【{CODE}】二手書8本一套" and 10 <= len(s["title"]) <= 60  # set of 8 second-hand books
    check_all(cfg)
    assert refresh_products(cfg)[0]["status"] == "ready"
    # a book priced in another currency: Shopee sells in NT$
    store.update_rows(cfg.csv_path, lambda rows: rows[2].update(currency="USD"))
    s = refresh_products(cfg)[0]
    assert s["status"] == "draft" and "3 (USD)" in s["errors"] and s["price"] == ""


def test_description_lists_every_book_by_number(cfg):
    text = refresh_products(cfg)[0]["description"]
    assert text.startswith(f"【{CODE}】二手書 8 本一套出售")
    assert (
        "#1 狄更斯的耶誕頌\n作者：瑪麗．波．奧斯本\n出版社：小天下\n書況：近全新\n價格：NT$100（新書定價 NT$250）"
        in text
    )
    assert "#8 Google必修的圖表簡報術：Google總監首度公開，教你做對圖表" in text
    assert "整套售價：NT$800" in text and "恕不拆賣" in text  # the footer you can edit
    assert (cfg.listings_dir / "description_footer.txt").exists()
    long = [
        {
            "position": str(i),
            "title": "長" * 200,
            "author": "作" * 100,
            "publisher": "出" * 100,
            "price": "100",
            "condition": "good",
        }
        for i in range(1, 31)
    ]
    assert len(set_description_zh("X", long, cfg, 3000, 3000.0)) <= 3000
    assert len(set_title_zh(CODE, [{"title": "很" * 80}] * 8)) <= 60


def test_discount_and_your_own_text_are_kept(cfg):
    cfg.set_discount = Config().set_discount  # the default: 20% off the whole set (8 折)
    check_all(cfg)
    s = refresh_products(cfg)[0]
    assert (s["original_price"], s["discount"], s["price"]) == ("800", "20", "640")
    assert "整套售價：NT$640（8 本單價合計 NT$800，整套 8 折）" in s["description"]
    s = save_product_fields(cfg, CODE, {"title": f"【{CODE}】神奇樹屋 8本", "discount": "15%"})
    assert s["title"] == f"【{CODE}】神奇樹屋 8本" and s["custom"] == "discount,title"
    assert s["price"] == "680" and "85 折" in s["description"]  # this set: 15% off
    store.update_rows(cfg.csv_path, lambda rows: rows[0].update(price="300"))
    s = refresh_products(cfg)[0]
    assert s["title"] == f"【{CODE}】神奇樹屋 8本"  # yours: kept
    assert (s["original_price"], s["price"]) == ("1000", "850")  # follows the books' prices
    s = save_product_fields(cfg, CODE, {"title": "", "discount": ""})  # empty = automatic again
    assert s["title"] == f"【{CODE}】二手書8本一套" and s["custom"] == "" and s["price"] == "800"
    for bad in ("abc", "100"):
        with pytest.raises(ValueError):
            save_product_fields(cfg, CODE, {"discount": bad})


def test_set_price_and_discount_follow_each_other(cfg):
    cfg.set_discount = Config().set_discount
    check_all(cfg)
    refresh_products(cfg)
    s = save_product_fields(cfg, CODE, {"price": "NT$600"})  # you type the price: the discount follows
    assert (s["price"], s["discount"], s["custom"]) == ("600", "25", "price")
    assert "75 折" in s["description"]
    store.update_rows(cfg.csv_path, lambda rows: rows[0].update(price="300"))
    s = refresh_products(cfg)[0]  # the books dearer: your price is kept
    assert (s["original_price"], s["price"], s["discount"]) == ("1000", "600", "40")
    s = save_product_fields(cfg, CODE, {"discount": "10"})  # now the discount: the price follows
    assert (s["price"], s["custom"]) == ("900", "discount")
    s = save_product_fields(cfg, CODE, {"price": "1,200"})  # dearer than the books: no discount
    assert (s["price"], s["discount"]) == ("1200", "0")
    s = save_product_fields(cfg, CODE, {"price": "", "discount": ""})
    assert (s["price"], s["discount"], s["custom"]) == ("800", "20", "")
    for bad in ("abc", "0", "-5"):
        with pytest.raises(ValueError):
            save_product_fields(cfg, CODE, {"price": bad})


def test_upload_file_is_shopees_template_filled_in(cfg):
    names = make_template(cfg.template_path)
    res = shopee.export(cfg)
    assert res["file"] is None and res["drafts"] and res["drafts"][0][0] == CODE  # book 4 not checked
    check_all(cfg)
    res = shopee.export(cfg)
    assert res["made"] == [CODE] and res["file"].exists()
    ws = openpyxl.load_workbook(res["file"])["Template"]
    row = {names[j]: ws.cell(6, j + 1).value for j in range(len(names))}
    assert ws.cell(5, 1).value == "格式"  # the explanation rows are left alone
    assert row["分類代碼"] == 101573  # found in the category list (其他書籍)
    assert row["商品名稱"].startswith(f"【{CODE}】") and row["主商品貨號"] == CODE
    assert row["價格"] == 800 and row["庫存"] == 1 and row["最低購買數量"] == 1  # both 1 unless you change them
    assert row["重量"] is None and row["長度"] is None and row["寬度"] is None and row["高度"] is None
    desc = the_set(cfg)["description"]
    assert row["商品描述"] == desc and desc.count("\n") > 40 and "#8 " in desc  # the whole text, every line
    assert ws.cell(6, names.index("商品描述") + 1).alignment.wrap_text  # Excel shows all of it
    assert row["7-ELEVEN"] == "開啟" and row["全家"] == "開啟" and row["黑貓宅急便"] == "關閉"
    assert row["規格名稱1"] is None and row["主商品圖片"] is None  # no variations; no photo links set up
    photos = sorted((cfg.listings_dir / "photos" / CODE).glob("*.jpg"))
    assert [p.name for p in photos] == [f"{i:02d}.jpg" for i in range(1, 10)]  # 1 set photo + 8 books
    assert all(Image.open(p).size == (1500, 1500) and p.stat().st_size < 2_000_000 for p in photos)
    s, li = the_set(cfg), the_listing(cfg)
    assert s["status"] == "listed" and len(s["images"].split(";")) == 9
    assert {b["status"] for b in books(cfg).values()} == {"listed"}
    assert li["platform"] == "shopee" and li["status"] == "exported" and li["export_file"] == str(res["file"])
    assert li["shopee_sku"] == CODE and li["shopee_category_id"] == "101573" and li["shopee_condition"] == "二手"
    assert li["price"] == "800" and li["title"] == s["title"]
    # run again: nothing new; --redo puts the set in a new file again
    assert shopee.export(cfg)["file"] is None
    again = shopee.export(cfg, redo=[CODE])
    assert again["made"] == [CODE] and again["file"] != res["file"]


def test_category_is_optional(cfg):
    check_all(cfg)
    names = make_template(cfg.template_path, category=False)
    res = shopee.export(cfg)  # no category found: Shopee will suggest one
    assert res["made"] == [CODE] and any("No category given" in n for n in res["notes"])
    ws = openpyxl.load_workbook(res["file"])["Template"]
    assert ws.cell(6, names.index("分類代碼") + 1).value is None and ws.cell(6, 2).value.startswith(f"【{CODE}】")
    assert the_listing(cfg)["shopee_category_id"] == ""


def test_without_template_nothing_is_made(cfg):
    check_all(cfg)
    cfg.shopee_category_id = "100644"
    with pytest.raises(shopee.ShopeeError, match="No Shopee template"):
        shopee.export(cfg)
    assert the_set(cfg)["status"] == "ready" and not store.read_table(cfg.listings_path)  # nothing changed
    assert not list(cfg.uploads_dir.glob("shopee_upload_*")) and not (cfg.listings_dir / "photos" / CODE).exists()
    assert {b["status"] for b in books(cfg).values()} == {"validated"}
    from pipeline.apps import editing

    assert editing.shopee_upload(cfg)["text"].startswith("Not done: No Shopee template")


def test_your_own_description_goes_in_the_file(cfg):
    names = make_template(cfg.template_path)
    check_all(cfg)
    refresh_products(cfg)
    mine = "【" + CODE + "】我自己寫的介紹\n第二行\n" + "很長的一段文字。" * 100
    save_product_fields(cfg, CODE, {"description": mine})
    res = shopee.export(cfg)
    ws = openpyxl.load_workbook(res["file"])["Template"]
    assert ws.cell(6, names.index("商品描述") + 1).value == mine


def order_report(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "訂單編號",
            "訂單狀態",
            "不成立原因",
            "退貨 / 退款狀態",
            "買家帳號",
            "訂單成立日期",
            "商品總價",
            "成交手續費",
            "金流與系統處理費",
            "包裹查詢號碼",
            "主商品貨號",
            "商品名稱",
            "訂單完成時間",
        ]
    )
    for r in rows:
        ws.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_orders_reserve_then_sell_the_set(cfg):
    make_template(cfg.template_path)
    check_all(cfg)
    shopee.export(cfg)
    title = the_set(cfg)["title"]
    # someone orders it
    order_report(
        cfg.shopee_import / "Order.all.20261012.xlsx",
        [["2610120ABC", "待出貨", "", "", "buyer1", "2026-10-12 10:00", "800", "", "", "", CODE, title, ""]],
    )
    out = shopee.run(cfg, do_export=False)
    assert any("ordered on Shopee" in t for t in out["text"])
    s, li = the_set(cfg), the_listing(cfg)
    assert s["status"] == "reserved" and li["shopee_order_sn"] == "2610120ABC" and li["status"] == "live"
    assert {b["status"] for b in books(cfg).values()} == {"listed"}
    assert not list(cfg.shopee_import.glob("*.xlsx")) and len(list((cfg.shopee_import / "done").iterdir())) == 1
    # the order is completed - found by the code in the title even without the SKU column
    order_report(
        cfg.shopee_import / "Order.completed.xlsx",
        [
            [
                "2610120ABC",
                "已完成",
                "",
                "",
                "buyer1",
                "2026-10-12 10:00",
                "800",
                "48",
                "16",
                "TW123",
                "",
                title,
                "2026-10-16 18:00",
            ]
        ],
    )
    shopee.run(cfg, do_export=False)
    s, li = the_set(cfg), the_listing(cfg)
    assert s["status"] == "sold" and s["sold_on"] == "shopee" and s["sold_price"] == "800"
    assert s["sold_at"] == "2026-10-16 18:00"
    assert li["status"] == "sold" and li["shopee_fees"] == "64" and li["shopee_tracking_no"] == "TW123"
    assert {b["status"] for b in books(cfg).values()} == {"sold"}


def test_cancelled_order_puts_the_set_back_on_sale(cfg):
    make_template(cfg.template_path)
    check_all(cfg)
    shopee.export(cfg)
    f = cfg.data_dir / "r.xlsx"
    order_report(f, [["A1", "待付款", "", "", "b", "", "800", "", "", "", CODE, "", ""]])
    shopee.sync(cfg, [f])
    assert the_set(cfg)["status"] == "reserved"
    order_report(f, [["A1", "不成立", "買家取消", "", "b", "", "800", "", "", "", CODE, "", ""]])
    res = shopee.sync(cfg, [f])
    assert the_set(cfg)["status"] == "listed" and "cancelled" in res["lines"][0]
    assert the_listing(cfg)["shopee_order_sn"] == ""


def test_product_list_and_pasted_link(cfg):
    make_template(cfg.template_path)
    check_all(cfg)
    shopee.export(cfg)
    cfg.shopee_shop_id = "55555"
    p = cfg.data_dir / "shopee_product_list.csv"  # Shopee's own product list (not our products.csv)
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["商品ID", "商品名稱", "主商品貨號", "狀態"])
        w.writerow(["24680", "whatever", CODE, "上架中"])
    shopee.sync(cfg, [p])
    li = the_listing(cfg)
    assert li["status"] == "live" and li["shopee_item_id"] == "24680"
    assert li["url"] == "https://shopee.tw/product/55555/24680"
    assert shopee.parse_url("https://shopee.tw/神奇樹屋-i.123.456?sp_atk=x") == ("123", "456")
    li = shopee.set_link(cfg, CODE, "https://shopee.tw/product/123/999")
    assert li["shopee_item_id"] == "999" and li["url"].endswith("/999")
    with pytest.raises(shopee.ShopeeError):
        shopee.set_link(cfg, CODE, "https://example.com")


def test_set_status_by_hand_and_takedown_elsewhere(cfg):
    check_all(cfg)
    refresh_products(cfg)
    store.write_table(
        cfg.listings_path,
        [{"product_code": CODE, "platform": "facebook", "status": "live", "url": "https://fb.example/1"}],
        store.LISTING_COLUMNS,
    )
    msgs = change_product_status(cfg, CODE, "sold", "sold at the school fair")
    assert msgs == [f"{CODE} is sold: take it down on facebook (https://fb.example/1)"]
    s = the_set(cfg)
    assert s["status"] == "sold" and s["sold_on"] == "manual" and s["sold_price"] == "800" and "fair" in s["notes"]
    assert {b["status"] for b in books(cfg).values()} == {"sold"}
    change_product_status(cfg, CODE, "ready")
    assert the_set(cfg)["status"] == "ready" and {b["status"] for b in books(cfg).values()} == {"validated"}
    with pytest.raises(ValueError):
        change_product_status(cfg, CODE, "reserved")


def test_books_marked_listed_by_hand_before_sets_existed(cfg):
    store.update_rows(cfg.csv_path, lambda rows: [r.update(status="listed") for r in rows])
    s = refresh_products(cfg)[0]
    assert s["status"] == "listed" and s["price"] == "800"
    assert shopee.export(cfg)["made"] == []  # already on sale: not listed twice


def test_old_books_csv_loses_the_unused_platform_columns(tmp_path):
    p = tmp_path / "books.csv"
    p.write_text("sku,trademe_id,ebay_id,fb_status,notes\na,1,2,live,n\n", encoding="utf-8")
    store.write_rows(p, store.read_rows(p))
    header = p.read_text(encoding="utf-8-sig").splitlines()[0]
    assert "trademe_id" not in header and "fb_status" not in header and "notes" in header


def test_password_protected_report(cfg, monkeypatch):
    import sys
    import types

    make_template(cfg.template_path)
    check_all(cfg)
    shopee.export(cfg)
    plain = cfg.data_dir / "plain.xlsx"
    order_report(plain, [["A9", "運送中", "", "", "b", "", "800", "", "", "", CODE, "", ""]])
    locked = cfg.shopee_import / "Order.toship.xlsx"
    locked.write_bytes(b"not a zip: an encrypted Office file")
    out = shopee.run(cfg, do_export=False)
    assert any("SHOPEE_EXPORT_PASSWORD" in t for t in out["text"]) and locked.exists()  # left there to retry
    fake = types.ModuleType("msoffcrypto")

    class OfficeFile:
        def __init__(self, f):
            pass

        def load_key(self, password):
            assert password == "123456"

        def decrypt(self, buf):
            buf.write(plain.read_bytes())

    fake.OfficeFile = OfficeFile
    monkeypatch.setitem(sys.modules, "msoffcrypto", fake)
    cfg.shopee_export_password = "123456"
    shopee.run(cfg, do_export=False)
    assert the_set(cfg)["status"] == "reserved" and not locked.exists()


def test_template_folder_any_name_and_uploads_folder(cfg):
    check_all(cfg)
    cfg.ensure_dirs()
    for d in (cfg.template_dir, cfg.uploads_dir, cfg.shopee_import, cfg.photos_dir, cfg.inbox):
        assert d.is_dir()  # nothing to make by hand
    make_template(cfg.template_dir / "a_first.xlsx")
    (cfg.template_dir / "~$a_first.xlsx").write_bytes(b"excel lock file")  # ignored
    (cfg.template_dir / "z_second.xlsx").write_bytes(b"not used: the first one is")
    assert shopee.find_template(cfg).name == "a_first.xlsx"
    res = shopee.export(cfg)
    assert res["file"].parent == cfg.uploads_dir and res["made"] == [CODE]


def test_photos_go_online_once(cfg):
    from pipeline.apps import photohost

    make_template(cfg.template_path)
    check_all(cfg)
    cfg.imgbb_api_key = "k3y"
    sent = []

    class Resp:
        status_code = 200

        def __init__(self, n):
            self.n = n

        def json(self):
            return {"success": True, "data": {"url": f"https://i.ibb.co/x{self.n}/{self.n:02d}.jpg"}}

    class Http:
        def post(self, url, data=None, timeout=None):
            assert url == photohost.IMGBB_URL and "k3y" not in url and data["key"] == "k3y"  # key not in the link
            sent.append(data["name"])
            return Resp(len(sent))

    res = shopee.export(cfg, http=Http())
    assert len(sent) == 9 and sent[0] == f"{CODE}-01"
    ws = openpyxl.load_workbook(res["file"])["Template"]
    names = [ws.cell(2, j).value for j in range(1, ws.max_column + 1)]
    assert ws.cell(6, names.index("主商品圖片") + 1).value == "https://i.ibb.co/x1/01.jpg"
    assert the_listing(cfg)["errors"] == ""
    shopee.export(cfg, redo=[CODE], http=Http())  # same photos: no new uploads
    assert len(sent) == 9

    class Down:
        def post(self, url, data=None, timeout=None):
            raise ConnectionError("offline")

    (cfg.photos_dir / CODE / "links.json").unlink()
    res = shopee.export(cfg, redo=[CODE], http=Down())
    assert res["file"] and any("photos not put online" in n for n in res["notes"])


def test_stock_and_minimum_purchase_quantity(cfg):
    names = make_template(cfg.template_path)
    check_all(cfg)
    p = refresh_products(cfg)[0]
    assert (p["stock"], p["min_purchase"]) == ("1", "1")
    p = save_product_fields(cfg, CODE, {"stock": "3", "min_purchase": "2"})
    assert (p["stock"], p["min_purchase"], p["custom"]) == ("3", "2", "")
    for bad in ({"min_purchase": "4"}, {"stock": "0"}, {"stock": "abc"}, {"min_purchase": "1000000"}):
        with pytest.raises(ValueError):  # min > stock: nobody could order it
            save_product_fields(cfg, CODE, bad)
    assert (the_set(cfg)["stock"], the_set(cfg)["min_purchase"]) == ("3", "2")  # nothing saved
    res = shopee.export(cfg)
    ws = openpyxl.load_workbook(res["file"])["Template"]
    row = {names[j]: ws.cell(6, j + 1).value for j in range(len(names))}
    assert row["庫存"] == 3 and row["最低購買數量"] == 2
    p = save_product_fields(cfg, CODE, {"stock": "", "min_purchase": ""})  # empty = 1 again
    assert (p["stock"], p["min_purchase"]) == ("1", "1")


def test_a_set_sold_per_book(cfg):
    from pipeline.apps.products import book_title_zh, set_sell_mode

    names = make_template(cfg.template_path)
    refresh_products(cfg)
    got = set_sell_mode(cfg, "20261008-153012", "S01", "book")  # each book its own product
    assert [p["product_code"] for p in got] == [f"{CODE}-{i:02d}" for i in range(1, 9)]
    prods = {p["product_code"]: p for p in refresh_products(cfg)}
    assert CODE not in prods  # the set itself is not for sale now
    one = prods[f"{CODE}-01"]
    assert one["kind"] == "book" and one["book_count"] == "1" and one["title"] == "狄更斯的耶誕頌｜瑪麗．波．奧斯本"
    assert (one["price"], one["discount"], one["original_price"], one["stock"]) == ("100", "0", "100", "1")
    d = one["description"]
    assert d.startswith("狄更斯的耶誕頌\n作者：瑪麗．波．奧斯本\n出版社：小天下") and "#1" not in d and "整套" not in d
    assert "書況：近全新" in d and "售價：NT$100（新書定價 NT$250）" in d and f"ISBN：{TW_ISBNS[0]}" in d
    assert prods[f"{CODE}-04"]["status"] == "draft" and "not checked" in prods[f"{CODE}-04"]["errors"]
    assert one["status"] == "ready"  # each book on its own: no waiting for book 4
    assert len(book_title_zh({"title": "雪"})) >= 10  # Shopee: at least 10 characters
    res = shopee.export(cfg)
    assert len(res["made"]) == 7 and f"{CODE}-04" not in res["made"]
    ws = openpyxl.load_workbook(res["file"])["Template"]
    rows = [{names[j]: ws.cell(r, j + 1).value for j in range(len(names))} for r in range(6, 13)]
    first = next(r for r in rows if r["主商品貨號"] == f"{CODE}-01")
    assert first["商品名稱"] == one["title"] and first["價格"] == 100 and first["商品描述"] == one["description"]
    photos = sorted((cfg.photos_dir / f"{CODE}-01").glob("*.jpg"))
    assert [p.name for p in photos] == ["01.jpg", "02.jpg"]  # its front (cover), its back
    assert books(cfg)["1"]["status"] == "listed" and books(cfg)["4"]["status"] == "to_check"
    with pytest.raises(ValueError, match="listed"):  # books on Shopee: can't go back to a set
        set_sell_mode(cfg, "20261008-153012", "S01", "set")


def test_back_to_one_product(cfg):
    from pipeline.apps.products import set_sell_mode

    check_all(cfg)
    refresh_products(cfg)
    save_product_fields(cfg, CODE, {"title": "【X】我的標題"})
    set_sell_mode(cfg, "20261008-153012", "S01", "book")
    set_sell_mode(cfg, "20261008-153012", "S01", "set")
    prods = refresh_products(cfg)
    assert [p["product_code"] for p in prods] == [CODE] and prods[0]["status"] == "ready"
    change_product_status(cfg, CODE, "listed")
    with pytest.raises(ValueError, match="the set is"):  # the set is on sale: stays a set
        set_sell_mode(cfg, "20261008-153012", "S01", "book")
    assert {b["sell_as"] for b in books(cfg).values()} == {""}
